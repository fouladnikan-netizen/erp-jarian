#!/usr/bin/env python3
"""Generate (and optionally import) the Product Master bulk-import template.

The workbook matches the live taxonomy + attribute schemas. Filling invented
group/category/type names will be rejected by bulk import.

  python3 backend/scripts/product-import-template.py generate
  python3 backend/scripts/product-import-template.py import --file path.xlsx
  python3 backend/scripts/product-import-template.py import --file path.xlsx --apply
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import shutil
import sys
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
API_BASE = os.environ.get("JARIAN_API_BASE", "http://127.0.0.1:3100")
USERNAME = os.environ.get("JARIAN_ADMIN_USER", "admin")
PASSWORD = os.environ.get("JARIAN_ADMIN_PASSWORD", "Admin123!")

FIXED_COLUMNS = [
    ("groupName", "گروه کالا — عین برگه انواع_کالا"),
    ("categoryName", "دسته کالا — عین برگه انواع_کالا"),
    ("typeName", "نوع کالا — عین برگه انواع_کالا"),
    ("displayNameOverride", "نام نمایشی اختیاری؛ خالی = نام ساخته‌شده سیستم"),
    ("brandName", "نام برند از برگه برندها؛ خالی مجاز"),
    ("baseUomCode", "کد واحد شمارش مثل PIECE یا KG — برگه واحدها"),
    ("salesUomCode", "کد واحد فروش"),
    ("purchaseUomCode", "کد واحد خرید"),
    ("unitWeight", "وزن واحد (عدد مثبت، کیلوگرم)"),
    ("customLengthAllowed", "TRUE یا FALSE"),
    ("weightProfileType", "FIXED | PER_LENGTH | DIMENSIONAL | MANUAL_ACTUAL"),
]


def api(method: str, path: str, token: str | None = None, body=None):
    data = None if body is None else json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        f"{API_BASE}{path}",
        data=data,
        method=method,
        headers={"Content-Type": "application/json"},
    )
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=60) as res:
            raw = res.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as err:
        detail = err.read().decode("utf-8", errors="replace")
        raise SystemExit(f"API {method} {path} failed ({err.code}): {detail}") from err
    except urllib.error.URLError as err:
        raise SystemExit(
            f"API در {API_BASE} در دسترس نیست. ابتدا npm run server را اجرا کنید. ({err.reason})"
        ) from err


def login() -> str:
    payload = api("POST", "/api/v1/auth/login", body={"username": USERNAME, "password": PASSWORD})
    token = payload.get("token") or payload.get("accessToken")
    if not token:
        raise SystemExit("Login did not return an access token.")
    return token


def is_transaction(binding: dict) -> bool:
    return binding.get("valueScope") == "TRANSACTION" or binding.get("attributeRole") == "TRANSACTION_ONLY"


def is_identity(definition: dict, binding: dict) -> bool:
    if is_transaction(binding):
        return False
    return definition.get("dataType") in {"DECIMAL", "INTEGER"}


def enum_values(entry: dict) -> list[str]:
    binding = entry["binding"]
    definition = entry["definition"]
    items = binding.get("effectiveAllowedValues") or definition.get("allowedValues") or []
    return [str(item.get("value")) for item in items if item.get("value") is not None]


def sample_value(entry: dict) -> str | None:
    binding = entry["binding"]
    definition = entry["definition"]
    if is_transaction(binding) or not binding.get("isRequired"):
        return None
    code = definition.get("code")
    data_type = definition.get("dataType")
    if data_type == "ENUM":
        values = enum_values(entry)
        return values[0] if values else None
    if code == "dimensions":
        return "20×20"
    if code == "grade":
        return "A2"
    if code == "size":
        return "8"
    if data_type in {"DECIMAL", "INTEGER"}:
        return "10"
    if data_type == "BOOLEAN":
        return "TRUE"
    return "نمونه"


def load_catalog(token: str):
    groups = api("GET", "/api/v1/product-taxonomy/groups", token)["items"]
    categories = api("GET", "/api/v1/product-taxonomy/categories", token)["items"]
    types = api("GET", "/api/v1/product-taxonomy/types", token)["items"]
    brands = api("GET", "/api/v1/brands", token)["items"]
    uoms = api("GET", "/api/v1/uom", token)["items"]
    definitions = api("GET", "/api/v1/attribute-definitions", token)["items"]

    groups_by_id = {g["id"]: g for g in groups}
    cats_by_id = {c["id"]: c for c in categories}

    schemas = {}

    def fetch_schema(type_id: str):
        payload = api("GET", f"/api/v1/attribute-definitions/schema/{type_id}", token)
        return type_id, payload.get("schema") or []

    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = [pool.submit(fetch_schema, t["id"]) for t in types]
        for fut in as_completed(futures):
            type_id, schema = fut.result()
            schemas[type_id] = schema

    type_rows = []
    product_attr_codes = {}
    for t in types:
        category = cats_by_id.get(t["categoryId"])
        group = groups_by_id.get(category["groupId"]) if category else None
        schema = schemas.get(t["id"]) or []
        required, optional, identity, txn, enum_hints = [], [], [], [], []
        for entry in schema:
            definition = entry["definition"]
            binding = entry["binding"]
            code = definition["code"]
            label = f"{code} ({definition.get('nameFa')})"
            if is_transaction(binding):
                txn.append(label)
                continue
            product_attr_codes[code] = definition
            if binding.get("isRequired"):
                required.append(label)
            else:
                optional.append(label)
            if is_identity(definition, binding):
                identity.append(code)
            if definition.get("dataType") == "ENUM":
                enum_hints.append(f"{code}:{'|'.join(enum_values(entry))}")
        type_rows.append({
            "groupName": group["name"] if group else "",
            "categoryName": category["name"] if category else "",
            "typeName": t["name"],
            "required": "، ".join(required),
            "optional": "، ".join(optional),
            "identity": "، ".join(identity),
            "enums": " ؛ ".join(enum_hints),
            "transaction": "، ".join(txn),
            "identityEmpty": "فقط یک کالا در این نوع (بدون ویژگی عددی هویت)" if not identity else "",
            "schema": schema,
        })

    type_rows.sort(key=lambda r: (r["groupName"], r["categoryName"], r["typeName"]))
    attr_codes = sorted(product_attr_codes)
    return {
        "groups": groups,
        "categories": categories,
        "types": types,
        "brands": brands,
        "uoms": uoms,
        "definitions": definitions,
        "type_rows": type_rows,
        "product_attr_codes": attr_codes,
        "product_attr_defs": product_attr_codes,
        "headers": [col[0] for col in FIXED_COLUMNS] + [f"attr:{code}" for code in attr_codes],
        "header_help": [col[1] for col in FIXED_COLUMNS] + [
            f"{product_attr_codes[code].get('nameFa')} / {product_attr_codes[code].get('dataType')}"
            for code in attr_codes
        ],
    }


PREFERRED_EXAMPLES = [
    ("مقاطع فولادی", "میلگرد", "میلگرد آجدار"),
    ("مقاطع فولادی", "پروفیل", "پروفیل مبلی"),
    ("مقاطع فولادی", "تیرآهن", "تیرآهن هاش"),
    ("مقاطع فولادی", "ورق سرد", "ورق گالوانیزه"),
    ("پیچ و مهره", "پیچ", None),
    ("استنلس استیل", None, None),
]


def row_from_type(catalog: dict, item: dict) -> dict:
    row = {key: "" for key in catalog["headers"]}
    row["groupName"] = item["groupName"]
    row["categoryName"] = item["categoryName"]
    row["typeName"] = item["typeName"]
    row["weightProfileType"] = "MANUAL_ACTUAL"
    for entry in item["schema"]:
        value = sample_value(entry)
        if value is None:
            continue
        row[f"attr:{entry['definition']['code']}"] = value
    return row


def build_example_rows(catalog: dict) -> list[dict]:
    by_key = {
        (item["groupName"], item["categoryName"], item["typeName"]): item
        for item in catalog["type_rows"]
    }
    rows = []
    used = set()

    def add_item(item):
        key = (item["groupName"], item["categoryName"], item["typeName"])
        if key in used:
            return
        used.add(key)
        rows.append(row_from_type(catalog, item))

    for group, category, type_name in PREFERRED_EXAMPLES:
        match = None
        if type_name:
            match = by_key.get((group, category, type_name))
        if not match:
            candidates = [
                item for item in catalog["type_rows"]
                if item["groupName"] == group
                and (not category or item["categoryName"] == category)
                and item["required"]
            ]
            match = candidates[0] if candidates else None
        if match:
            add_item(match)

    for item in catalog["type_rows"]:
        if item["groupName"] not in {r["groupName"] for r in rows} and item["required"]:
            add_item(item)
        if len(rows) >= 8:
            break
    return rows


def csv_text(catalog: dict) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(catalog["headers"])
    writer.writerow(["#" + catalog["header_help"][0], *catalog["header_help"][1:]])
    for example in build_example_rows(catalog):
        writer.writerow([example.get(h, "") for h in catalog["headers"]])
    return buf.getvalue()


def require_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        return {
            "Workbook": Workbook,
            "Comment": Comment,
            "Alignment": Alignment,
            "Font": Font,
            "PatternFill": PatternFill,
            "get_column_letter": get_column_letter,
            "DataValidation": DataValidation,
        }
    except ImportError:
        raise SystemExit("openpyxl لازم است: python3 -m pip install openpyxl")


def autosize(ws, max_width=48):
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col[:40]:
            value = "" if cell.value is None else str(cell.value)
            width = min(max_width, max(width, min(len(value) + 2, max_width)))
        ws.column_dimensions[letter].width = width


def write_xlsx(path: Path, catalog: dict):
    ox = require_openpyxl()
    wb = ox["Workbook"]()
    header_fill = ox["PatternFill"]("solid", fgColor="1F4E79")
    header_font = ox["Font"](color="FFFFFF", bold=True)
    wrap = ox["Alignment"](wrap_text=True, vertical="top", horizontal="right")

    guide = wb.active
    guide.title = "راهنما"
    guide.sheet_view.rightToLeft = True
    guide_lines = [
        "تمپلیت ورود کالا — جریان",
        "فقط برگه «محصولات» را پر کنید. نام گروه / دسته / نوع را از برگه «انواع_کالا» کپی کنید؛ نام جعلی رد می‌شود.",
        "نام کالا را سیستم از نوع + ویژگی‌ها می‌سازد. اگر نام تجاری جدا می‌خواهید displayNameOverride را پر کنید.",
        "ویژگی‌ها ستون‌های attr:کد لاتین هستند. فقط ویژگی‌های همان نوع را پر کنید؛ بقیه را خالی بگذارید.",
        "ویژگی الزامی هر نوع در ستون «ویژگی‌های الزامی» برگه انواع_کالا آمده است.",
        "هویت کالا = ویژگی‌های عددی سطح محصول (مثل size و thickness). مقدار یکسان = همان کالا و ردیف تکراری رد می‌شود.",
        "تیرآهن هاش: یک نوع است. سبک/سنگین را در attr:kind با light یا heavy بگذارید — HEA/HEB نوع جدا نیستند.",
        "استنلس: گرید معمولاً خودِ نوع کالاست (مثلاً ورق استیل 304). ستون grade را فقط اگر در انواع_کالا برای همان نوع آمده پر کنید.",
        "سایز لوله و schedule باید عدد باشند (مثلاً 0.5) نه «1/2 اینچ».",
        "واحد را با کد لاتین بنویسید (PIECE، KG، SHEET، METER) نه نام فارسی.",
        "weightProfileType فقط یکی از: FIXED ، PER_LENGTH ، DIMENSIONAL ، MANUAL_ACTUAL",
        "customLengthAllowed فقط TRUE یا FALSE",
        "ENUM را دقیقاً از برگه مقادیر_ENUM وارد کنید.",
        "پس از تکمیل: برگه محصولات را Save As → CSV UTF-8 کنید و در ویترین → ورود دسته‌ای بارگذاری کنید. یا همین فایل xlsx را با دستور import همین اسکریپت بفرستید.",
        "سقف هر ارسال: ۲۰۰۰ ردیف.",
    ]
    for idx, line in enumerate(guide_lines, start=1):
        guide.cell(idx, 1, line)
        guide.cell(idx, 1).alignment = wrap
    guide.column_dimensions["A"].width = 110

    products = wb.create_sheet("محصولات")
    products.sheet_view.rightToLeft = True
    products.freeze_panes = "A2"
    for col, header in enumerate(catalog["headers"], start=1):
        cell = products.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.comment = ox["Comment"](catalog["header_help"][col - 1], "Jarian")
    for r_idx in range(2, 82):
        for c_idx in range(1, len(catalog["headers"]) + 1):
            products.cell(r_idx, c_idx, None)
    last_col = ox["get_column_letter"](len(catalog["headers"]))
    uom_last = max(2, len(catalog["uoms"]) + 1)
    brand_last = max(2, len([b for b in catalog["brands"] if b.get("isActive", True)]) + 1)
    group_last = max(2, len(catalog["groups"]) + 1)
    for formula, cols in (
        ('"FIXED,PER_LENGTH,DIMENSIONAL,MANUAL_ACTUAL"', [11]),
        ('"TRUE,FALSE"', [10]),
        (f"'واحدها'!$A$2:$A${uom_last}", [6, 7, 8]),
        (f"'برندها'!$A$2:$A${brand_last}", [5]),
        (f"'فهرست_گروه‌ها'!$A$2:$A${group_last}", [1]),
    ):
        dv = ox["DataValidation"](type="list", formula1=formula, allow_blank=True)
        products.add_data_validation(dv)
        for col in cols:
            letter = ox["get_column_letter"](col)
            dv.add(f"{letter}2:{letter}2000")
    autosize(products)
    products.auto_filter.ref = f"A1:{last_col}80"

    sample = wb.create_sheet("نمونه")
    sample.sheet_view.rightToLeft = True
    for col, header in enumerate(catalog["headers"], start=1):
        cell = sample.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
    for r_idx, example in enumerate(build_example_rows(catalog), start=2):
        for c_idx, header in enumerate(catalog["headers"], start=1):
            sample.cell(r_idx, c_idx, example.get(header) or None)
    autosize(sample)

    types_ws = wb.create_sheet("انواع_کالا")
    types_ws.sheet_view.rightToLeft = True
    type_headers = [
        "groupName", "categoryName", "typeName",
        "ویژگی‌های الزامی", "ویژگی‌های اختیاری", "هویت عددی",
        "مقادیر ENUM", "ویژگی تراکنش (پر نکنید)", "یادداشت هویت",
    ]
    for col, header in enumerate(type_headers, start=1):
        cell = types_ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
    for r_idx, item in enumerate(catalog["type_rows"], start=2):
        types_ws.cell(r_idx, 1, item["groupName"])
        types_ws.cell(r_idx, 2, item["categoryName"])
        types_ws.cell(r_idx, 3, item["typeName"])
        types_ws.cell(r_idx, 4, item["required"])
        types_ws.cell(r_idx, 5, item["optional"])
        types_ws.cell(r_idx, 6, item["identity"])
        types_ws.cell(r_idx, 7, item["enums"])
        types_ws.cell(r_idx, 8, item["transaction"])
        types_ws.cell(r_idx, 9, item["identityEmpty"])
    autosize(types_ws, 60)
    types_ws.auto_filter.ref = f"A1:I{len(catalog['type_rows']) + 1}"

    attrs = wb.create_sheet("ویژگی‌ها")
    attrs.sheet_view.rightToLeft = True
    for col, header in enumerate(["code", "nameFa", "dataType", "ستون تمپلیت"], start=1):
        cell = attrs.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
    for r_idx, code in enumerate(catalog["product_attr_codes"], start=2):
        definition = catalog["product_attr_defs"][code]
        attrs.cell(r_idx, 1, code)
        attrs.cell(r_idx, 2, definition.get("nameFa"))
        attrs.cell(r_idx, 3, definition.get("dataType"))
        attrs.cell(r_idx, 4, f"attr:{code}")
    autosize(attrs)

    enums = wb.create_sheet("مقادیر_ENUM")
    enums.sheet_view.rightToLeft = True
    for col, header in enumerate(["groupName", "categoryName", "typeName", "attrCode", "value"], start=1):
        cell = enums.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
    r_idx = 2
    for item in catalog["type_rows"]:
        for entry in item["schema"]:
            if is_transaction(entry["binding"]) or entry["definition"].get("dataType") != "ENUM":
                continue
            for value in enum_values(entry):
                enums.cell(r_idx, 1, item["groupName"])
                enums.cell(r_idx, 2, item["categoryName"])
                enums.cell(r_idx, 3, item["typeName"])
                enums.cell(r_idx, 4, entry["definition"]["code"])
                enums.cell(r_idx, 5, value)
                r_idx += 1
    autosize(enums)
    if r_idx > 2:
        enums.auto_filter.ref = f"A1:E{r_idx - 1}"

    uoms = wb.create_sheet("واحدها")
    uoms.sheet_view.rightToLeft = True
    for col, header in enumerate(["code", "nameFa"], start=1):
        cell = uoms.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
    for r_idx, uom in enumerate(catalog["uoms"], start=2):
        uoms.cell(r_idx, 1, uom.get("code"))
        uoms.cell(r_idx, 2, uom.get("nameFa"))
    autosize(uoms)

    brands = wb.create_sheet("برندها")
    brands.sheet_view.rightToLeft = True
    brands.cell(1, 1, "brandName").fill = header_fill
    brands.cell(1, 1).font = header_font
    active_brands = [b for b in catalog["brands"] if b.get("isActive", True)]
    for r_idx, brand in enumerate(active_brands, start=2):
        brands.cell(r_idx, 1, brand.get("brandName") or brand.get("name"))
    autosize(brands)

    groups = wb.create_sheet("فهرست_گروه‌ها")
    groups.sheet_view.rightToLeft = True
    groups.cell(1, 1, "groupName").fill = header_fill
    groups.cell(1, 1).font = header_font
    for r_idx, group in enumerate(catalog["groups"], start=2):
        groups.cell(r_idx, 1, group.get("name"))
    autosize(groups)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_outputs(catalog: dict) -> list[Path]:
    docs_dir = ROOT / "Docs" / "templates"
    public_dir = ROOT / "public" / "templates"
    downloads = Path.home() / "Downloads"
    xlsx_name = "jarian-product-import.xlsx"
    csv_name = "jarian-product-import.csv"
    primary_xlsx = docs_dir / xlsx_name
    write_xlsx(primary_xlsx, catalog)
    csv_body = csv_text(catalog)
    written = [primary_xlsx]
    for directory in (docs_dir, public_dir, downloads):
        directory.mkdir(parents=True, exist_ok=True)
        csv_path = directory / csv_name
        csv_path.write_text("\ufeff" + csv_body, encoding="utf-8")
        written.append(csv_path)
        xlsx_path = directory / xlsx_name
        if xlsx_path != primary_xlsx:
            shutil.copy2(primary_xlsx, xlsx_path)
            written.append(xlsx_path)
    return written


def row_from_cells(headers: list[str], cells: list[str]) -> dict | None:
    row = {"attributes": {}}
    for idx, header in enumerate(headers):
        value = cells[idx].strip() if idx < len(cells) and cells[idx] is not None else ""
        if value is None:
            continue
        value = str(value).strip()
        if not value:
            continue
        if header.startswith("attr:"):
            row["attributes"][header[5:]] = value
        else:
            row[header] = value
    if not (row.get("groupName") and row.get("categoryName") and row.get("typeName")):
        return None
    if str(row["groupName"]).startswith("#"):
        return None
    return row


def read_template_rows(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.excel_tab if "\t" in text.splitlines()[0] else csv.excel
        reader = csv.reader(io.StringIO(text), dialect=dialect)
        rows = list(reader)
        headers = [h.strip() for h in rows[0]]
        return [parsed for cells in rows[1:] if (parsed := row_from_cells(headers, cells))]
    ox = require_openpyxl()
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb["محصولات"] if "محصولات" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    parsed = []
    for cells in rows[1:]:
        values = ["" if c is None else str(c) for c in cells]
        item = row_from_cells(headers, values)
        if item:
            parsed.append(item)
    return parsed


def cmd_generate() -> None:
    token = login()
    catalog = load_catalog(token)
    written = write_outputs(catalog)
    print(f"types={len(catalog['type_rows'])} attr_columns={len(catalog['product_attr_codes'])}")
    for path in written:
        print(path)


def cmd_import(file_path: Path, apply: bool) -> None:
    rows = read_template_rows(file_path)
    if not rows:
        raise SystemExit("هیچ ردیف کالایی در فایل نیست.")
    if len(rows) > 2000:
        raise SystemExit("سقف هر ارسال ۲۰۰۰ ردیف است.")
    token = login()
    mode = "APPLY" if apply else "DRY_RUN"
    result = api("POST", "/api/v1/products/bulk-import", token, {"mode": mode, "rows": rows})
    batch = result.get("batch") or result
    print(json.dumps({
        "mode": batch.get("mode"),
        "id": batch.get("id"),
        "totalRows": batch.get("totalRows"),
        "acceptedRows": batch.get("acceptedRows"),
        "rejectedRows": batch.get("rejectedRows"),
        "rowResults": batch.get("rowResults"),
    }, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description="Product import template generate/import")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("generate")
    import_p = sub.add_parser("import")
    import_p.add_argument("--file", required=True, type=Path)
    import_p.add_argument("--apply", action="store_true", help="Persist products. Default is DRY_RUN.")
    args = parser.parse_args()
    if args.command == "generate":
        cmd_generate()
    else:
        cmd_import(args.file, args.apply)


if __name__ == "__main__":
    main()
